"""Rank-change workbook: 2018 paper ranks vs 2026 update ranks (common-country basis), with live formulas."""
import os
import csv, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation

W = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))) + '/'
C = ["USA","GBR","DEU","SWE","FRA","NLD","CHE","DNK","CAN","JPN","AUS"]
colors = json.load(open(W+'spec/country_colors.json'))['colors']; LAB = {c: colors[c]['label'] for c in C}
spec = list(csv.DictReader(open(W+'spec/indicators.csv'))); S = {r['indicator_id']: r for r in spec}
R = {(r['indicator_id'], r['country']): r for r in csv.DictReader(open(W+'outputs/data/master_modern_dataset.csv'))}
FIGNAME = {'1':'Spending','4':'Population health','5':'Workforce & capacity','7':'Utilization','9':'Pharmaceuticals','10':'Access & quality','11':'Equity'}

# ---------- desirable direction: +1 higher is better, -1 lower is better, 0 no normative direction
DIR = {}
for i in S: DIR[i] = (0, 'Descriptive / context: no normative better-or-worse direction')
def setd(ids, d, why):
    for i in ids: DIR[i] = (d, why)
setd(['F1_05'], -1, 'Less poverty is better')
setd(['F1_17'], 1, 'Wider insurance coverage is better')
setd(['F4_01','F4_02','F4_03'], -1, 'Lower risk-factor prevalence is better')
setd(['F4_04','F4_05','F4_06','F4_07'], 1, 'Longer life expectancy is better')
setd(['F4_08','F4_09','F4_10','F4_11','F4_12'], -1, 'Lower mortality / low birth weight is better')
setd(['F9_13'], -1, 'Lower antibiotic use is generally better (stewardship)')
setd(['F10_01','F10_03'], 1, 'Better access/experience')
setd(['F10_02'], -1, 'Fewer long specialist waits is better')
setd(['F10_07','F10_08'], 1, 'Higher prevention coverage is better')
setd(['F10_09','F10_10','F10_11','F10_12'], -1, 'Lower adverse outcome rate is better')
setd(['F10_13','F10_14','F10_15','F10_16'], -1, 'Fewer avoidable admissions is better')
setd(['F10_04'], 1, 'Higher satisfaction is better'); setd(['F10_05','F10_06'], -1, 'Less dissatisfaction is better')
setd(['F11_01'], -1, 'Lower pro-rich inequity is better (0 = equal use by need)')
setd(['F11_02','F11_03'], -1, 'Lower out-of-pocket burden is better (financial protection)')
setd(['F11_04','F11_05','F11_06'], -1, 'Less cost-related unmet need is better')

# ---------- comparability (include flag + note)
EXCL_IND = {'F5_02': 'Definition narrowed in 2026 (OECD categories exclude general internists) and affects countries unequally',
            'F5_03': 'Same as primary-care share',
            'F11_09': '2018 used a patchwork of national urban/rural definitions; 2026 uses OECD TL3 typology',
            'F11_10': '2018 used a patchwork of national urban/rural definitions; 2026 uses OECD TL3 typology'}
EXCL_CTY = {('F1_07','CHE'): 'Mandatory insurance now classed as private in WHO GHED', ('F1_07','NLD'): 'Mandatory insurance now classed as private in WHO GHED',
            ('F7_02','FRA'): 'Series break 2016 (psychiatric stays added)', ('F5_04','DNK'): 'Series break 2022 in nurse definition',
            ('F4_03','DEU'): '2026 value self-reported vs measured in 2018',
            ('F11_07','NLD'): 'Rural series revised (definition change)', ('F11_07','CHE'): 'Rural series revised (definition change)'}
for i in ['F1_09','F1_10','F1_11','F1_12','F1_13','F1_14','F1_15','F1_16']:
    EXCL_CTY[(i,'NLD')] = 'Health-accounts series break 2021 (reclassification between functions)'
CAUT_IND = {'F1_04': 'Land area now vs surface area in 2018 (Canada/US order is an artifact)',
            'F1_08': '2018 basis uncertain (stated nominal US$, text says PPP); 2026 is PPP',
            'F4_05': 'WHO HALE revision changed levels; 2026 values are 2021 estimates',
            'F4_08': '2018 values partly model-based (GBD); 2026 registry-based',
            'F5_01': '2018 from mixed national sources; 2026 OECD practising',
            'F5_05': 'Mixed OECD/national sources in both years', 'F5_06': 'Mixed OECD/national sources in both years',
            'F5_09': 'Ratio uses matched-year wages in 2026', 'F5_10': 'Ratio uses matched-year wages in 2026', 'F5_11': 'Ratio uses matched-year wages in 2026',
            'F5_16': 'OECD now counts beds in institutions only',
            'F9_11': 'Market basis varies by country; 2018 Netherlands value looks anomalous', 'F9_12': 'Market basis varies by country',
            'F9_13': '2026 community sector only', 'F10_01': '2023 survey moved partly online', 'F10_02': '2023 survey moved partly online',
            'F10_14': 'Prevalence denominator from IDF Atlas 2024', 'F11_01': 'Different study/method (OECD 2019 vs Health at a Glance 2013)',
            'F11_02': 'Denominator now current (not total) health expenditure'}

num = lambda v: None if v in ('NA', '', None) else float(v)
def fig_rank(i, which):
    vals = {c: num(S[i][f'orig_{c}']) if which == 'old' else num(R[(i, c)]['value_plot']) for c in C}
    order = sorted([c for c in C if vals[c] is not None], key=lambda c: (-vals[c], C.index(c)))
    return {c: order.index(c) + 1 for c in order}, len(order)

long = []
for s in spec:
    i = s['indicator_id']; r_old, n_old = fig_rank(i, 'old'); r_new, n_new = fig_rank(i, 'new')
    for c in C:
        o, n = num(s[f'orig_{c}']), num(R[(i, c)]['value_plot'])
        inc, note = 1, CAUT_IND.get(i, '')
        if o is None or n is None: inc, note = 0, 'No value in ' + ('both years' if o is None and n is None else '2018' if o is None else '2026')
        elif i in EXCL_IND: inc, note = 0, 'Excluded: ' + EXCL_IND[i]
        elif (i, c) in EXCL_CTY: inc, note = 0, 'Excluded: ' + EXCL_CTY[(i, c)]
        elif i.startswith('F7_') and c == 'USA' and R[(i, c)]['obs_year'] and int(R[(i, c)]['obs_year']) <= 2010:
            note = (note + '; ' if note else '') + 'US OECD data still 2009-10 (same vintage as 2018), so US moves reflect other countries'
        elif i.startswith('F1_') and i[3:] >= '09' and i <= 'F1_16' and c == 'USA':
            note = (note + '; ' if note else '') + 'US 2018 health-accounts data (latest split); 2018 paper adjusted US inpatient +2 pts'
        if (i, c) in (('F10_15', 'USA'), ('F10_13', 'USA')):
            note = (note + '; ' if note else '') + 'US moved from ICD-9 to ICD-10-CM coding in Oct 2015; part of the change may be a coding effect'
        long.append(dict(fig=s['figure'], dom=FIGNAME[s['figure']], sec=s['section'], id=i, lab=s['label'], c=c, cl=LAB[c],
                         o=o, n=n, y=R[(i, c)]['obs_year'] if n is not None else '', fr_o=r_old.get(c), fr_n=r_new.get(c), inc=inc, note=note))

wb = Workbook()
F = 'Arial'; fnt = lambda **k: Font(name=F, size=k.pop('size', 10), **k)
HDR = PatternFill('solid', fgColor='1F3A5F'); WH = fnt(bold=True, color='FFFFFF')
BLUE = fnt(color='0000FF'); GREEN = PatternFill('solid', fgColor='C6EFCE'); RED = PatternFill('solid', fgColor='FFC7CE'); GREY = PatternFill('solid', fgColor='EDEDED')
UPF = PatternFill('solid', fgColor='DDEBF7'); DNF = PatternFill('solid', fgColor='FCE4D6')
def header(ws, row, vals, height=30):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row, j, v); c.font = WH; c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = height

# ---------- README
ws = wb.active; ws.title = 'README'
lines = [("Rank changes: 2018 JAMA paper (data mostly 2013–2016) vs 2026 update (data mostly 2023–2025)", True, 13),
 ("How to read", True, 11),
 ("• Rank convention is the paper's: rank 1 = HIGHEST value in the row, whether high is good or bad.", False, 10),
 ("• Rank change = 2018 rank − 2026 rank. Positive (▲) = the country moved toward the top (its value is now higher RELATIVE to the others); negative (▼) = moved down.", False, 10),
 ("• To keep the comparison fair, ranks are recomputed among only the countries that have a comparable value in BOTH years (column 'n compared'). The ranks printed in the figures are shown alongside for reference.", False, 10),
 ("• Performance change = rank change × desirable direction (Directions sheet): +1 if higher is better (e.g. life expectancy), −1 if lower is better (e.g. infant mortality), 0 if the indicator is descriptive (spending, utilization, workforce). Positive = the country now does BETTER relative to the other countries; negative = WORSE.", False, 10),
 ("• This is relative standing only: a country can improve in absolute terms yet fall in rank if others improved faster.", False, 10),
 ("Sheets", True, 11),
 ("Country_Summary — per country: how many comparable indicators improved, worsened or did not move in relative terms, overall and by domain; with chart.", False, 10),
 ("Performance_Matrix — indicator × country grid of performance change (green better, red worse) for indicators with a desirable direction.", False, 10),
 ("Rank_Movement_Matrix — indicator × country grid of raw rank change for ALL indicators (blue up, orange down), including descriptive ones.", False, 10),
 ("Rank_Changes — the underlying table (one row per indicator × country), with values, ranks, and comparability notes. All ranks and changes are live formulas.", False, 10),
 ("Directions — the desirable direction assigned to each indicator (blue = editable). These are judgment calls; change any of them and every sheet updates.", False, 10),
 ("Comparability rules", True, 11),
 ("Excluded (Include = 0): no value in one year; indicators whose definition changed in a way that affects countries unequally (primary-care share, urban/rural physicians); country values hit by a series break (e.g. Netherlands health accounts 2021, France mental-health discharges, Denmark nurses) or a definition switch (Germany self-reported weight; Swiss/Dutch public spending). Include flags are editable.", False, 10),
 ("Caution notes flag rows that remain in the comparison but where some movement may come from method changes (e.g. survey mode, data revisions). See the Comparability note column.", False, 10),
 ("Indicators with no 2026 data (drug prices, NCEs, health-system views, etc.) cannot be compared and show as excluded.", False, 10)]
for t, b, sz in lines:
    ws.append([t]); c = ws.cell(ws.max_row, 1); c.font = fnt(bold=b, size=sz); c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 160

# ---------- Directions (inputs)
wd = wb.create_sheet('Directions')
header(wd, 1, ['Indicator ID', 'Figure', 'Indicator', 'Desirable direction (+1 higher better, −1 lower better, 0 none)', 'Rationale'])
dv = DataValidation(type='list', formula1='"1,-1,0"', allow_blank=False); wd.add_data_validation(dv)
for s in spec:
    i = s['indicator_id']; wd.append([i, int(s['figure']), s['label'], DIR[i][0], DIR[i][1]])
    cell = wd.cell(wd.max_row, 4); cell.font = BLUE; cell.alignment = Alignment(horizontal='center'); dv.add(cell)
for j, w in enumerate([10, 7, 50, 20, 60], 1): wd.column_dimensions[L(j)].width = w
for row in wd.iter_rows(min_row=2):
    for c in row:
        if c.column != 4: c.font = fnt()
wd.freeze_panes = 'A2'
ND = len(spec) + 1

# ---------- Rank_Changes (long, formulas)
wl = wb.create_sheet('Rank_Changes')
H = ['Figure', 'Domain', 'Indicator ID', 'Indicator', 'Country', 'Country label', '2018 value (paper)', '2026 value', '2026 obs. year',
     'Rank in 2018 figure', 'Rank in 2026 figure', 'Include (1/0)', 'n compared', 'Rank 2018 (common set)', 'Rank 2026 (common set)',
     'Rank change (+ = up)', 'Desirable direction', 'Performance change (+ = better)', 'Relative performance', 'Comparability note', 'Key']
header(wl, 1, H, 45)
N = len(long) + 1
for k, d in enumerate(long, 2):
    wl.append([int(d['fig']), d['dom'], d['id'], d['lab'], d['c'], d['cl'], d['o'] if d['o'] is not None else 'NA', d['n'] if d['n'] is not None else 'NA',
               d['y'], d['fr_o'] or 'NA', d['fr_n'] or 'NA', d['inc'],
               f'=IF(L{k}=1,COUNTIFS($C$2:$C${N},C{k},$L$2:$L${N},1),"")',
               f'=IF(L{k}=1,1+COUNTIFS($C$2:$C${N},C{k},$L$2:$L${N},1,$G$2:$G${N},">"&G{k}),"")',
               f'=IF(L{k}=1,1+COUNTIFS($C$2:$C${N},C{k},$L$2:$L${N},1,$H$2:$H${N},">"&H{k}),"")',
               f'=IF(L{k}=1,N{k}-O{k},"")',
               f'=INDEX(Directions!$D$2:$D${ND},MATCH(C{k},Directions!$A$2:$A${ND},0))',
               f'=IF(L{k}=1,P{k}*Q{k},"")',
               f'=IF(L{k}<>1,"Excluded",IF(Q{k}=0,"No direction",IF(R{k}>0,"Better",IF(R{k}<0,"Worse","No change"))))',
               d['note'], f'=C{k}&"|"&E{k}'])
    wl.cell(k, 12).font = BLUE
for row in wl.iter_rows(min_row=2):
    for c in row:
        if c.column != 12: c.font = fnt()
for j, w in enumerate([6, 16, 8, 40, 6, 9, 9, 9, 8, 8, 8, 8, 8, 9, 9, 9, 8, 11, 13, 60, 12], 1): wl.column_dimensions[L(j)].width = w
wl.freeze_panes = 'G2'; wl.auto_filter.ref = f'A1:U{N}'
wl.conditional_formatting.add(f'S2:S{N}', CellIsRule(operator='equal', formula=['"Better"'], fill=GREEN))
wl.conditional_formatting.add(f'S2:S{N}', CellIsRule(operator='equal', formula=['"Worse"'], fill=RED))
wl.conditional_formatting.add(f'P2:P{N}', CellIsRule(operator='greaterThan', formula=['0'], fill=UPF))
wl.conditional_formatting.add(f'P2:P{N}', CellIsRule(operator='lessThan', formula=['0'], fill=DNF))
wl.column_dimensions['U'].hidden = True

# ---------- matrices
def matrix(name, col, ids, title, up, dn, sub):
    wm = wb.create_sheet(name); wm.sheet_view.showGridLines = False
    wm['A1'] = title; wm['A1'].font = fnt(bold=True, size=13)
    wm['A2'] = sub; wm['A2'].font = fnt(size=9, italic=True)
    header(wm, 4, ['Indicator ID', 'Domain', 'Indicator'] + [LAB[c] for c in C], 20)
    r = 5
    for i in ids:
        s = S[i]; wm.cell(r, 1, i); wm.cell(r, 2, FIGNAME[s['figure']]); wm.cell(r, 3, s['label'])
        for j, c in enumerate(C, 4):
            cell = wm.cell(r, j, f'=IFERROR(IF(INDEX(Rank_Changes!${col}$2:${col}${N},MATCH($A{r}&"|"&"{c}",Rank_Changes!$U$2:$U${N},0))="","–",'
                                 f'INDEX(Rank_Changes!${col}$2:${col}${N},MATCH($A{r}&"|"&"{c}",Rank_Changes!$U$2:$U${N},0))),"–")')
            cell.alignment = Alignment(horizontal='center'); cell.number_format = '+0;-0;0'; cell.font = fnt()
        for j in (1, 2, 3): wm.cell(r, j).font = fnt()
        r += 1
    rng = f'D5:N{r-1}'
    wm.conditional_formatting.add(rng, FormulaRule(formula=['AND(ISNUMBER(D5),D5>0)'], fill=up))
    wm.conditional_formatting.add(rng, FormulaRule(formula=['AND(ISNUMBER(D5),D5<0)'], fill=dn))
    wm.conditional_formatting.add(rng, FormulaRule(formula=['D5="–"'], fill=GREY))
    for j, w in enumerate([9, 18, 46] + [9] * 11, 1): wm.column_dimensions[L(j)].width = w
    wm.freeze_panes = 'D5'
    return wm
dir_ids = [s['indicator_id'] for s in spec if DIR[s['indicator_id']][0] != 0]
matrix('Performance_Matrix', 'R', dir_ids, 'Relative performance change, 2018 → 2026 (positions gained among comparable countries)', GREEN, RED,
       'Green = better relative standing; red = worse; 0 = same position; – = not comparable. Indicators with a desirable direction only (see Directions). Example: +2 for a mortality row means the country moved 2 places toward the low-mortality end.')
matrix('Rank_Movement_Matrix', 'P', [s['indicator_id'] for s in spec], 'Raw rank movement, 2018 → 2026 (all indicators; + = moved toward rank 1 / higher value)', UPF, DNF,
       'Blue = moved up (higher value relative to others); orange = moved down; – = not comparable. No judgment of good or bad.')

# ---------- Country summary
wsu = wb.create_sheet('Country_Summary', 1); wsu.sheet_view.showGridLines = False
wsu['A1'] = 'Relative performance by country, 2018 paper vs 2026 update'; wsu['A1'].font = fnt(bold=True, size=13)
wsu['A2'] = 'Counts of comparable indicators with a desirable direction where the country\'s rank (among countries with data in both years) improved, worsened or did not change.'; wsu['A2'].font = fnt(size=9, italic=True)
header(wsu, 4, ['Country', 'Indicators compared', 'Better', 'Worse', 'No change', 'Net (better − worse)', 'Share better', 'Mean positions gained'], 32)
rg = lambda col: f'Rank_Changes!${col}$2:${col}${N}'
for k, c in enumerate(C, 5):
    wsu.append([LAB[c], f'=COUNTIFS({rg("E")},"{c}",{rg("S")},"Better")+COUNTIFS({rg("E")},"{c}",{rg("S")},"Worse")+COUNTIFS({rg("E")},"{c}",{rg("S")},"No change")',
                f'=COUNTIFS({rg("E")},"{c}",{rg("S")},"Better")', f'=COUNTIFS({rg("E")},"{c}",{rg("S")},"Worse")', f'=COUNTIFS({rg("E")},"{c}",{rg("S")},"No change")',
                f'=C{k}-D{k}', f'=IF(B{k}=0,"",C{k}/B{k})',
                f'=IFERROR(AVERAGEIFS({rg("R")},{rg("E")},"{c}",{rg("Q")},"<>0",{rg("L")},1),"")'])
    for j in range(1, 9): wsu.cell(k, j).font = fnt(bold=(c == 'USA'))
    wsu.cell(k, 1).fill = PatternFill('solid', fgColor=colors[c]['hex'].lstrip('#'))
    wsu.cell(k, 7).number_format = '0%'; wsu.cell(k, 8).number_format = '+0.00;-0.00;0.00'
last = 4 + len(C)
wsu.conditional_formatting.add(f'F5:F{last}', CellIsRule(operator='greaterThan', formula=['0'], fill=GREEN))
wsu.conditional_formatting.add(f'F5:F{last}', CellIsRule(operator='lessThan', formula=['0'], fill=RED))
# by domain
r0 = last + 3
wsu.cell(r0 - 1, 1, 'Net (better − worse) by domain').font = fnt(bold=True, size=11)
doms = [('4', 'Population health'), ('10', 'Access & quality'), ('11', 'Equity'), ('1', 'Spending'), ('9', 'Pharmaceuticals')]
header(wsu, r0, ['Country'] + [d for _, d in doms], 30)
for k, c in enumerate(C, r0 + 1):
    wsu.cell(k, 1, LAB[c]).font = fnt(bold=(c == 'USA'))
    wsu.cell(k, 1).fill = PatternFill('solid', fgColor=colors[c]['hex'].lstrip('#'))
    for j, (f, _) in enumerate(doms, 2):
        cell = wsu.cell(k, j, f'=COUNTIFS({rg("E")},"{c}",{rg("A")},{f},{rg("S")},"Better")-COUNTIFS({rg("E")},"{c}",{rg("A")},{f},{rg("S")},"Worse")')
        cell.font = fnt(bold=(c == 'USA')); cell.alignment = Alignment(horizontal='center'); cell.number_format = '+0;-0;0'
rl = r0 + len(C)
wsu.conditional_formatting.add(f'B{r0+1}:F{rl}', CellIsRule(operator='greaterThan', formula=['0'], fill=GREEN))
wsu.conditional_formatting.add(f'B{r0+1}:F{rl}', CellIsRule(operator='lessThan', formula=['0'], fill=RED))
wsu.cell(rl + 1, 1, 'Domains with directional indicators only (Spending: poverty and coverage; Pharmaceuticals: antibiotic use). Workforce and Utilization indicators are descriptive — see Rank_Movement_Matrix.').font = fnt(size=9, italic=True)
for j, w in enumerate([14, 12, 9, 9, 10, 12, 10, 12], 1): wsu.column_dimensions[L(j)].width = w
ch = BarChart(); ch.type = 'bar'; ch.style = 10; ch.title = 'Net change in relative standing (indicators better minus worse)'
ch.add_data(Reference(wsu, min_col=6, min_row=4, max_row=last), titles_from_data=True); ch.set_categories(Reference(wsu, min_col=1, min_row=5, max_row=last))
ch.legend = None; ch.height = 9; ch.width = 15; ch.y_axis.title = 'Net count'; ch.x_axis.delete = False; ch.y_axis.delete = False; ch.x_axis.tickLblPos = 'low'
wsu.add_chart(ch, 'J4')
wb.calculation.fullCalcOnLoad = True  # Excel computes formulas on open
wb.save(W + 'outputs/Papanicolas2018_rank_changes_2018_vs_2026.xlsx'); print('saved', N - 1, 'rows')
