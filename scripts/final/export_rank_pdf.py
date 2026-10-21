"""Optional: export the two matrix sheets of the rank-change workbook to a 2-page PDF (needs LibreOffice)."""
import os, shutil, subprocess, sys, tempfile
import openpyxl
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
src = os.path.join(ROOT, 'outputs', 'Papanicolas2018_rank_changes_2018_vs_2026.xlsx')
dst = os.path.join(ROOT, 'outputs', 'Rank_change_matrices_2018_vs_2026.pdf')
soffice = shutil.which('soffice') or shutil.which('libreoffice')
if not soffice: sys.exit('LibreOffice not found; skipping PDF export')
tmp = tempfile.mkdtemp(); x = os.path.join(tmp, 'Rank_change_matrices.xlsx')
wb = openpyxl.load_workbook(src)
for n in wb.sheetnames:
    ws = wb[n]
    if n not in ('Performance_Matrix', 'Rank_Movement_Matrix'): ws.sheet_state = 'hidden'; continue
    ps = ws.page_setup; ps.paperSize = ws.PAPERSIZE_LETTER
    ps.orientation = 'landscape' if n == 'Performance_Matrix' else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ps.fitToWidth = 1; ps.fitToHeight = 1
    ws.page_margins.left = ws.page_margins.right = 0.35; ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.print_options.horizontalCentered = True; ws.print_area = f'A1:N{ws.max_row}'
    ws.oddFooter.center.text = '2018 JAMA paper (Papanicolas et al.) vs 2026 data update - rank changes among countries with comparable data in both years'
    ws.oddFooter.center.size = 7
wb.active = wb.sheetnames.index('Performance_Matrix'); wb.save(x)
subprocess.run([soffice, '--headless', '--convert-to', 'pdf', '--outdir', tmp, x], check=True, capture_output=True, timeout=300)
shutil.copy(os.path.join(tmp, 'Rank_change_matrices.pdf'), dst); print('wrote', dst)
