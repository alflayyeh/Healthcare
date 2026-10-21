import csv, json, glob, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo

C = ["USA","GBR","DEU","SWE","FRA","NLD","CHE","DNK","CAN","JPN","AUS"]
col = json.load(open('spec/country_colors.json'))['colors']
LAB = {c: col[c]['label'] for c in C}
spec = list(csv.DictReader(open('spec/indicators.csv'))); S = {r['indicator_id']: r for r in spec}
rows = list(csv.DictReader(open('outputs/data/master_modern_dataset.csv')))
R = {(r['indicator_id'], r['country']): r for r in rows}
audit = {}
for f in glob.glob('results/*_audit.csv'):
    for r in csv.DictReader(open(f)): audit[r['indicator_id']] = r
meta = json.load(open('spec/figure_meta_2026.json'))['figures']
qc = list(csv.DictReader(open('outputs/data/qc_results.csv')))
cpi = list(csv.DictReader(open('results/oecd_spend_us_cpi.csv')))
wages = list(csv.DictReader(open('results/oecd_spend_wages_current_ppp.csv')))

F = 'Arial'
def font(**k): return Font(name=F, size=k.pop('size', 10), **k)
HDR = PatternFill('solid', fgColor='1F3A5F'); WH = font(bold=True, color='FFFFFF')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
STATUS_FILL = {'DIRECTLY UPDATED': 'D9EAD3', 'CLOSE MODERN EQUIVALENT': 'FFF2CC', 'NOT COMPARABLE': 'F4CCCC', 'UNAVAILABLE': 'E6E6E6'}
wb = Workbook()

def table_sheet(ws, header, data, widths, name, wrap_cols=()):
    ws.append(header)
    for d in data: ws.append([("'" + x if False else x) for x in d])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.value = cell.value; cell.data_type = 's'
    for j, w in enumerate(widths, 1): ws.column_dimensions[L(j)].width = w
    for cell in ws[1]: cell.font = WH; cell.fill = HDR; cell.alignment = Alignment(wrap_text=True, vertical='center')
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = font(); cell.alignment = Alignment(wrap_text=cell.column in wrap_cols, vertical='top')
    ws.freeze_panes = 'B2'
    t = Table(displayName=name, ref=f"A1:{L(len(header))}{len(data)+1}")
    t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True); ws.add_table(t)

# overrides of agent audit by orchestrator decisions
OVR = {
 'F1_04': ('CLOSE MODERN EQUIVALENT', 'World Bank WDI AG.LND.TOTL.K2 (land area)', 'Land area excluding inland waters, thousand km²', 'Original plotted total surface area (incl. inland waters). The WB surface-area series now includes coastal waters for CAN/FRA/SWE (Canada 15 634 vs 9985), so land area (consistent for all 11, matching the row label) is plotted; surface area kept in value_as_retrieved.'),
 'F5_02': ('CLOSE MODERN EQUIVALENT', 'OECD Health Statistics DF_PHYS_CAT (% of physicians)', 'Generalist medical practitioners + general paediatricians as % of all physicians', 'Original used a functionality-based definition incl. general internal medicine & geriatrics from mixed national sources. OECD/Eurostat no longer separate general internists, so a single OECD definition is applied to all countries (narrower; levels lower for DEU, SWE, GBR, CHE, FRA, USA). Japan has no OECD generalist category → NA. National functional values (US KFF 48.7%, Japan MHLW 24.3%) kept in value_as_retrieved.'),
 'F5_03': ('CLOSE MODERN EQUIVALENT', 'OECD Health Statistics DF_PHYS_CAT', '100 − primary care share (same definition as F5_02)', 'See F5_02.'),
 'F5_09': ('CLOSE MODERN EQUIVALENT', 'Derived: OECD remuneration (US: Medscape) / OECD average annual wage', 'GP remuneration ÷ average annual wage, same country and year, current USD PPP', 'Original divided by average wage (2016). Now computed with matching observation years.'),
 'F5_10': ('CLOSE MODERN EQUIVALENT', 'Derived', 'Specialist remuneration ÷ average annual wage, same year', 'As F5_09.'),
 'F5_11': ('CLOSE MODERN EQUIVALENT', 'Derived', 'Hospital nurse remuneration ÷ average annual wage, same year', 'As F5_09.'),
}
def aud(i):
    a = dict(audit.get(i, {}))
    if i in OVR:
        a['status'], a['modern_source'], a['modern_definition'], extra = OVR[i]
        a['change_vs_2018'] = extra + (' | Agent note: ' + a['change_vs_2018'] if a.get('change_vs_2018') else '')
    return a

# ---------------- README
ws = wb.active; ws.title = 'README'
txt = [
 ("Health Care Spending in the US and Other High-Income Countries — 2026 data update", True),
 ("Modern rebuild of the ranked matrix figures (Figures 1, 4, 5, 7, 9, 10, 11) of Papanicolas I, Woskie LR, Jha AK. JAMA. 2018;319(10):1024-1039.", False),
 ("Data retrieved October 7, 2026. Each value is the newest comparable observation for that country; years differ by country and indicator (no interpolation or modelling).", False),
 ("", False),
 ("Sheets", True),
 ("Figure_Inventory — the 7 qualifying figures, their sections and indicators (figures 2, 3, 6, 8 are bar charts and the Table is text; excluded).", False),
 ("Source_Map — per indicator: original 2018 source → modern source, comparability status, years used, countries available, what changed (audit + change log).", False),
 ("Master_Data — one row per indicator × country (1089 rows): plotted value, value as retrieved, unit, observation year, 2018 value, source, dataset code, direct URL, flags, decisions, notes (source appendix).", False),
 ("Wide_Data — plotted values in country columns with live formulas for the Mean and the number of countries contributing (n).", False),
 ("Fig1 … Fig11 — the recreated figures as colored Excel matrices (ranked highest → lowest, Mean column linked to Wide_Data). High-resolution PNG/SVG/PDF versions are delivered separately.", False),
 ("Remuneration_Calc — CPI deflation of physician/nurse pay to 2025 USD and pay-to-wage ratios, as live formulas.", False),
 ("Change_Log — indicators whose modern definition or source differs materially from 2018.", False),
 ("Research_Gaps — indicators or countries that could not be updated reliably, and why.", False),
 ("QC — automated checks (rank order, duplicates, units, magnitudes vs 2018, sources, traceability to saved raw downloads).", False),
 ("", False),
 ("Status definitions", True),
 ("DIRECTLY UPDATED — same indicator and source still exist.  CLOSE MODERN EQUIVALENT — concept substantially the same; dataset/definition changed.  NOT COMPARABLE — modern data exist but are definitionally inconsistent (values kept in Master_Data, not plotted).  UNAVAILABLE — no credible current equivalent (row shown as NA).", False),
 ("", False),
 ("Key decisions", True),
 ("1. Source hierarchy of the 2018 paper kept: health spending % GDP / per capita from WHO GHED via World Bank (as in 2018), not OECD, even though OECD has later provisional years (OECD values are in agent_notes).", False),
 ("2. A value is plotted only if its definition matches the other countries in the row; excluded values remain in Master_Data with the reason (orchestrator_decision).", False),
 ("3. Remuneration converted to 2025 USD with US CPI-U, mirroring the paper's conversion to 2017 USD.", False),
 ("4. Ranking convention unchanged: every row is ranked highest → lowest regardless of whether high is good or bad.", False),
 ("5. Mean = unweighted mean of countries with data; denominator shown as n in Wide_Data and as (n = x) in the figures.", False),
]
for t, b in txt:
    ws.append([t]); ws.cell(ws.max_row, 1).font = font(bold=b, size=12 if b else 10)
    ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 150

# ---------------- Figure inventory
ws = wb.create_sheet('Figure_Inventory')
inv = []
for fig, title in [('1','Spending'),('4','Population Health'),('5','Workforce and Structural Capacity'),('7','Utilization'),('9','Pharmaceuticals'),('10','Access and Quality'),('11','Distribution and Equity')]:
    ids = [r for r in spec if r['figure'] == fig]
    secs = list(dict.fromkeys(r['section'] for r in ids if r['section']))
    inv.append([f'Figure {fig}', title, len(ids), ' · '.join(secs), '; '.join(r['label'] for r in ids)])
table_sheet(ws, ['Figure', 'Domain', 'No. of indicators', 'Sections', 'Indicators'], inv, [10, 30, 12, 60, 140], 'Inventory', wrap_cols=(4, 5))

# ---------------- Source map / audit
ws = wb.create_sheet('Source_Map')
sm = []
for s in spec:
    i = s['indicator_id']; a = aud(i)
    pl = [R[(i, c)] for c in C if R[(i, c)]['value_plot'] != 'NA']
    ys = sorted({r['obs_year'] for r in pl})
    sm.append([i, f"Figure {s['figure']}", s['section'], s['label'], s['definition'], s['orig_source'], s['orig_year'],
               a.get('modern_source', ''), a.get('modern_definition', ''), a.get('status', ''),
               (ys[0] + ('–' + ys[-1] if ys[-1] != ys[0] else '')) if ys else '', len(pl),
               ', '.join(LAB[c] for c in C if R[(i, c)]['value_plot'] == 'NA'), a.get('change_vs_2018', ''), a.get('notes', '')])
table_sheet(ws, ['Indicator ID', 'Figure', 'Section', 'Indicator (figure label)', '2018 definition', 'Original 2018 source', 'Original year',
                 'Modern source', 'Modern definition', 'Status', 'Years plotted', 'Countries plotted', 'NA countries', 'Change vs 2018', 'Notes'],
            sm, [9, 9, 22, 34, 40, 34, 10, 40, 40, 22, 11, 9, 24, 60, 60], 'SourceMap', wrap_cols=(3, 4, 5, 6, 8, 9, 13, 14, 15))
for r in range(2, len(sm) + 2):
    c = ws.cell(r, 10); c.fill = PatternFill('solid', fgColor=STATUS_FILL.get(c.value, 'FFFFFF'))

# ---------------- Master data
ws = wb.create_sheet('Master_Data')
hdr = ['Figure', 'Domain section', 'Indicator ID', 'Indicator', 'Definition (2018)', 'Country', 'Country label', 'Value plotted', 'Value as retrieved', 'Unit',
       'Observation year', '2018 value', 'Original 2018 source', 'Modern source', 'Dataset code', 'Direct source URL', 'Retrieval date', 'Source flags',
       'Status', 'Orchestrator decision', 'Notes', 'Research agent']
md = []
for r in rows:
    s = S[r['indicator_id']]
    num = lambda v: None if v in ('NA', '') else float(v)
    md.append([int(r['figure']), r['section'], r['indicator_id'], r['label'], s['definition'], r['country'], LAB[r['country']],
               num(r['value_plot']) if r['value_plot'] != 'NA' else 'NA', num(r['value_as_retrieved']) if r['value_as_retrieved'] != 'NA' else 'NA',
               r['unit'], r['obs_year'], num(r['orig_2018_value']) if r['orig_2018_value'] not in ('NA', '') else 'NA', s['orig_source'], r['modern_source'],
               r['dataset_code'], r['url'], r['retrieval_date'], r['flags'], aud(r['indicator_id']).get('status', r['status']), r['orchestrator_decision'],
               r['agent_notes'], r['research_agent']])
table_sheet(ws, hdr, md, [7, 20, 9, 32, 30, 8, 10, 11, 11, 22, 10, 9, 28, 32, 20, 40, 11, 14, 20, 40, 60, 12], 'MasterData')
ws.freeze_panes = 'H2'

# ---------------- Remuneration calc (formulas)
ws = wb.create_sheet('Remuneration_Calc')
ws.append(['Reference year for constant dollars', 2025]); ws['A1'].font = font(bold=True); ws['B1'].font = font(color='0000FF')
ws.append(['CPI-U (reference year)', '=INDEX(CPI_Wages!$B:$B,MATCH($B$1,CPI_Wages!$A:$A,0))'])
ws.append([])
h = ['Indicator ID', 'Country', 'Remuneration, current USD PPP (obs year)', 'Observation year', 'CPI-U obs year', 'Deflation factor', 'Remuneration, 2025 USD',
     'Average wage, current USD PPP (same year)', 'Ratio to average wage', 'Ratio indicator']
ws.append(h)
for c_ in ws[4]: c_.font = WH; c_.fill = HDR; c_.alignment = Alignment(wrap_text=True)
remrow = {}
for i, ri in (('F5_05', 'F5_09'), ('F5_06', 'F5_10'), ('F5_07', 'F5_11')):
    for c in C:
        r = R[(i, c)]
        if r['value_as_retrieved'] == 'NA': continue
        n = ws.max_row + 1; remrow[(i, c)] = n; remrow[(ri, c)] = n
        y = int(r['obs_year'])
        ratio_ok = not (c == 'JPN' and ri != 'F5_11')
        ws.append([i, c, float(r['value_as_retrieved']), y,
                   f'=INDEX(CPI_Wages!$B:$B,MATCH(D{n},CPI_Wages!$A:$A,0))', f'=$B$2/E{n}', f'=ROUND(C{n}*F{n},0)',
                   f'=IFERROR(INDEX(CPI_Wages!$F:$F,MATCH(B{n}&"|"&D{n},CPI_Wages!$E:$E,0)),"NA")',
                   f'=IFERROR(ROUND(C{n}/H{n},2),"NA")' if ratio_ok else 'NA (Japan combined figure)', ri])
        ws.cell(n, 3).font = font(color='0000FF'); ws.cell(n, 4).font = font(color='0000FF')
for j, w in enumerate([10, 8, 18, 10, 10, 10, 14, 18, 12, 10], 1): ws.column_dimensions[L(j)].width = w
ws.column_dimensions['A'].width = 34
ws.append([]); ws.append(['Blue = values retrieved (Master_Data "Value as retrieved"); black = formulas. CPI-U: BLS CUUR0000SA0 annual average. Wages: OECD average annual wages, current USD PPP.'])

ws = wb.create_sheet('CPI_Wages')
ws.append(['Year', 'CPI-U annual average', 'Source', '', 'Key (country|year)', 'Average annual wage, current USD PPP', 'Country', 'Year'])
for c_ in ws[1]: c_.font = WH; c_.fill = HDR
mx = max(len(cpi), len(wages))
for k in range(mx):
    a = [int(cpi[k]['year']), float(cpi[k]['cpi_u_annual_avg']), cpi[k]['source_url']] if k < len(cpi) else ['', '', '']
    b = [f"{wages[k]['country']}|{wages[k]['year']}", float(wages[k]['value_current_usd_ppp']), wages[k]['country'], int(wages[k]['year'])] if k < len(wages) else ['', '', '', '']
    ws.append(a + [''] + b)
for j, w in enumerate([8, 14, 60, 2, 14, 18, 8, 8], 1): ws.column_dimensions[L(j)].width = w
ws.append([]); ws.append(['Wage source: OECD Average annual wages (OECD.ELS.SAE DSD_EARNINGS@AV_AN_WAGE), national currency ÷ OECD PPP for private consumption; reproduces OECD USD PPP values.'])

# ---------------- Wide data with formulas
ws = wb.create_sheet('Wide_Data')
ws.append(['Indicator ID', 'Figure', 'Indicator', 'Unit'] + [LAB[c] for c in C] + ['Mean', 'n', 'Status'] + [f'Year {LAB[c]}' for c in C])
for c_ in ws[1]: c_.font = WH; c_.fill = HDR; c_.alignment = Alignment(wrap_text=True)
WROW = {}
for s in spec:
    i = s['indicator_id']; n = ws.max_row + 1; WROW[i] = n
    vals = []
    for k, c in enumerate(C):
        r = R[(i, c)]
        if r['value_plot'] == 'NA': vals.append('NA')
        elif i in ('F5_05', 'F5_06', 'F5_07'): vals.append(f"=Remuneration_Calc!G{remrow[(i, c)]}")
        elif i in ('F5_09', 'F5_10', 'F5_11'): vals.append(f"=Remuneration_Calc!I{remrow[(i, c)]}")
        else: vals.append(float(r['value_plot']))
    rng = f"E{n}:O{n}"
    ws.append([i, int(s['figure']), s['label'], R[(i, 'USA')]['unit'] or s['unit']] + vals +
              [f'=IF(COUNT({rng})=0,"NA",ROUND(AVERAGE({rng}),{max(int(s["decimals"]),1) if int(s["decimals"]) else 0}))', f'=COUNT({rng})', aud(i).get('status', '')] +
              [R[(i, c)]['obs_year'] if R[(i, c)]['value_plot'] != 'NA' else '' for c in C])
for row in ws.iter_rows(min_row=2):
    for c_ in row: c_.font = font()
for j, w in enumerate([9, 6, 40, 22] + [9] * 11 + [9, 4, 22] + [7] * 11, 1): ws.column_dimensions[L(j)].width = w
ws.freeze_panes = 'E2'

# ---------------- Figure sheets
for fig in ['1', '4', '5', '7', '9', '10', '11']:
    m = meta[fig]; ws = wb.create_sheet(f'Fig{fig}')
    ws.sheet_view.showGridLines = False
    ws['A1'] = m['title']; ws['A1'].font = font(bold=True, size=13)
    ws.append([]); hdr = ['Rank (highest to lowest)'] + [str(k) for k in range(1, 12)] + ['Mean']
    ws.append(hdr)
    for c_ in ws[3]: c_.font = font(bold=True); c_.border = BOX; c_.alignment = Alignment(horizontal='center')
    ws['A3'].alignment = Alignment(horizontal='left')
    sec = None
    for s in [x for x in spec if x['figure'] == fig]:
        i = s['indicator_id']
        if s['section'] and s['section'] != sec:
            sec = s['section']; ws.append([sec]); n = ws.max_row
            ws.merge_cells(f'A{n}:M{n}'); ws[f'A{n}'].font = font(bold=True); ws[f'A{n}'].fill = PatternFill('solid', fgColor='F2F2F2')
        elif not s['section']: sec = None
        lab = m.get('label_overrides', {}).get(i, s['label']).replace('\n', ' ')
        cells = sorted(C, key=lambda c: (R[(i, c)]['value_plot'] == 'NA', -float(R[(i, c)]['value_plot']) if R[(i, c)]['value_plot'] != 'NA' else 0, C.index(c)))
        d = int(s['decimals'])
        def fmt(v):
            if v == 'NA': return 'NA'
            x = float(v); t = f'{x:,.{d}f}' if abs(x) >= 10000 else f'{x:.{d}f}'
            return t.replace(',', ' ')
        ws.append([lab] + [f"{LAB[c]}\n{fmt(R[(i, c)]['value_plot'])}" for c in cells] + [f"=Wide_Data!P{WROW[i]}"])
        n = ws.max_row; ws.row_dimensions[n].height = 30
        ws.cell(n, 1).font = font(bold=not s['section']); ws.cell(n, 1).alignment = Alignment(wrap_text=True, vertical='top')
        for k, c in enumerate(cells, 2):
            cl = ws.cell(n, k); cl.fill = PatternFill('solid', fgColor=col[c]['hex'].lstrip('#'))
            cl.alignment = Alignment(wrap_text=True, vertical='top'); cl.font = font(size=9); cl.border = Border(left=Side(style='thin', color='FFFFFF'), right=Side(style='thin', color='FFFFFF'), top=Side(style='thin', color='FFFFFF'), bottom=Side(style='thin', color='FFFFFF'))
        ws.cell(n, 13).font = font(size=9); ws.cell(n, 13).alignment = Alignment(vertical='top', horizontal='left')
    ws.append([])
    notes = [m['general_note']] + [f"{f['marker']} {f['text']}" if isinstance(f, dict) else f for f in m.get('footnotes', [])] + [m.get('source_note', '')]
    for t in notes:
        ws.append([t]); n = ws.max_row; ws.merge_cells(f'A{n}:M{n}')
        ws[f'A{n}'].alignment = Alignment(wrap_text=True, vertical='top'); ws[f'A{n}'].font = font(size=8)
        ws.row_dimensions[n].height = max(14, 12 * (len(t) // 170 + 1))
    ws.column_dimensions['A'].width = 34
    for k in range(2, 13): ws.column_dimensions[L(k)].width = 10.5
    ws.column_dimensions['M'].width = 9

# ---------------- Change log
ws = wb.create_sheet('Change_Log')
cl = []
for s in spec:
    i = s['indicator_id']; a = aud(i)
    decs = sorted({R[(i, c)]['orchestrator_decision'] for c in C if R[(i, c)]['orchestrator_decision']})
    if a.get('status') != 'DIRECTLY UPDATED' or decs:
        cl.append([i, f"Figure {s['figure']}", s['label'], a.get('status', ''), s['orig_source'], a.get('modern_source', ''), a.get('change_vs_2018', ''),
                   ' | '.join(f"{LAB[c]}: {R[(i, c)]['orchestrator_decision']}" for c in C if R[(i, c)]['orchestrator_decision'])[:3000]])
table_sheet(ws, ['Indicator ID', 'Figure', 'Indicator', 'Status', '2018 source', 'Modern source', 'What changed vs 2018', 'Country-level decisions'],
            cl, [9, 9, 32, 22, 30, 34, 80, 80], 'ChangeLog', wrap_cols=(3, 5, 6, 7, 8))

# ---------------- Research gaps
ws = wb.create_sheet('Research_Gaps')
gp = []
for s in spec:
    i = s['indicator_id']; a = aud(i)
    na = [c for c in C if R[(i, c)]['value_plot'] == 'NA']
    if a.get('status') in ('UNAVAILABLE', 'NOT COMPARABLE'):
        gp.append([i, f"Figure {s['figure']}", s['label'], a['status'], 'All 11', a.get('change_vs_2018', '') + ' ' + a.get('notes', '')])
    elif na:
        reasons = '; '.join(f"{LAB[c]}: {(R[(i, c)]['orchestrator_decision'] or R[(i, c)]['agent_notes'] or 'not reported by source')[:220]}" for c in na)
        gp.append([i, f"Figure {s['figure']}", s['label'], 'Partial coverage', ', '.join(LAB[c] for c in na), reasons])
table_sheet(ws, ['Indicator ID', 'Figure', 'Indicator', 'Gap type', 'Countries affected', 'Why'], gp, [9, 9, 34, 18, 28, 140], 'Gaps', wrap_cols=(3, 5, 6))

# ---------------- QC
ws = wb.create_sheet('QC')
table_sheet(ws, ['Indicator ID', 'Test', 'Result', 'Detail / resolution'], [[q['indicator_id'], q['test'], q['result'], q['detail']] for q in qc], [10, 50, 10, 120], 'QCResults', wrap_cols=(4,))
wb.calculation.fullCalcOnLoad = True  # Excel computes formulas on open
wb.save('outputs/Papanicolas2018_modern_update.xlsx'); print('saved')
